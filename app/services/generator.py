"""Document generator service for NAO documents."""
import os
from datetime import datetime
from typing import Dict, List
from docx import Document
from jinja2 import Environment, FileSystemLoader, select_autoescape

from app.utils.docx_utils import (
    create_document,
    add_heading,
    add_styled_paragraph,
    add_bullet_list,
    add_table_with_data,
    add_signature_block,
    save_document,
)
from app.utils.legal_refs import format_legal_reference
from app.config import settings


# Setup Jinja2 environment for templates
template_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "templates", "documents")
jinja_env = Environment(
    loader=FileSystemLoader(template_dir),
    autoescape=select_autoescape()
)


def generate_letter_ouverture(campaign: Dict, output_path: str) -> str:
    """
    Generate opening letter for NAO negotiations.

    Args:
        campaign: Campaign data dictionary
        output_path: Output file path

    Returns:
        Path to generated document
    """
    doc = create_document()

    # Header
    add_styled_paragraph(doc, campaign.get("organization_name", "CGT"), bold=True)
    add_styled_paragraph(doc, datetime.now().strftime("%d/%m/%Y"))
    doc.add_paragraph()

    # Recipient
    add_styled_paragraph(doc, f"À l'attention de la Direction")
    add_styled_paragraph(doc, campaign.get("establishment_name", ""))
    doc.add_paragraph()

    # Subject
    add_styled_paragraph(
        doc,
        f"Objet : Ouverture des Négociations Annuelles Obligatoires {datetime.now().year}",
        bold=True,
    )
    doc.add_paragraph()

    # Body
    add_styled_paragraph(doc, "Madame, Monsieur,")
    doc.add_paragraph()

    add_styled_paragraph(
        doc,
        f"Conformément aux dispositions des articles L2242-1 et suivants du Code du travail, "
        f"nous souhaitons engager les Négociations Annuelles Obligatoires (NAO) pour l'année {datetime.now().year}."
    )
    doc.add_paragraph()

    add_styled_paragraph(
        doc,
        "Ces négociations porteront sur les thèmes suivants :"
    )

    themes = [
        "Rémunération, salaires effectifs et partage de la valeur ajoutée",
        "Temps de travail et organisation du travail",
        "Égalité professionnelle entre les femmes et les hommes",
        "Qualité de vie au travail et conditions de travail",
    ]
    add_bullet_list(doc, themes)
    doc.add_paragraph()

    add_styled_paragraph(
        doc,
        "Nous vous proposons de nous rencontrer dans les meilleurs délais afin de convenir "
        "d'un calendrier de négociation."
    )
    doc.add_paragraph()

    # Legal references
    add_styled_paragraph(doc, "Références légales :", bold=True)
    legal_refs = [
        format_legal_reference("L2242-1"),
        format_legal_reference("L2242-2"),
        format_legal_reference("L2242-5"),
    ]
    add_bullet_list(doc, legal_refs)
    doc.add_paragraph()

    # Closing
    add_styled_paragraph(
        doc,
        "Dans l'attente de votre réponse, nous vous prions d'agréer, Madame, Monsieur, "
        "l'expression de nos salutations distinguées."
    )
    doc.add_paragraph()

    # Signature
    add_signature_block(
        doc,
        campaign.get("contact_name", "Délégué·e syndical·e"),
        "Délégué·e syndical·e",
        campaign.get("organization_name", "CGT"),
    )

    save_document(doc, output_path)
    return output_path


def generate_revendications(campaign: Dict, output_path: str) -> str:
    """
    Generate detailed demands document.

    Args:
        campaign: Campaign data dictionary
        output_path: Output file path

    Returns:
        Path to generated document
    """
    doc = create_document()

    # Title
    add_heading(
        doc,
        f"Revendications NAO {datetime.now().year} - {campaign.get('establishment_name', '')}",
        level=1,
    )
    doc.add_paragraph()

    # Introduction
    add_styled_paragraph(
        doc,
        f"Dans le cadre des Négociations Annuelles Obligatoires (NAO) {datetime.now().year}, "
        f"les organisations syndicales présentent les revendications suivantes au bénéfice "
        f"des {campaign.get('employee_count', 'XX')} salarié·es de {campaign.get('establishment_name', '')}."
    )
    doc.add_paragraph()

    demands = campaign.get("demands", {})

    # 1. Salaires
    if demands.get("salaries", {}).get("enabled"):
        add_heading(doc, "1. Rémunération et Salaires", level=2)
        salary_demands = demands["salaries"]

        if salary_demands.get("general_increase"):
            add_styled_paragraph(
                doc,
                f"Augmentation générale des salaires : {salary_demands['general_increase']}%"
            )

        if salary_demands.get("seniority_bonus"):
            add_styled_paragraph(doc, "Mise en place d'une prime d'ancienneté")

        if salary_demands.get("salary_grid_update"):
            add_styled_paragraph(doc, "Mise à jour de la grille salariale")

        add_styled_paragraph(
            doc,
            f"Base légale : {format_legal_reference('L2242-2')}"
        )
        doc.add_paragraph()

    # 2. Congés
    if demands.get("leaves", {}).get("enabled"):
        add_heading(doc, "2. Congés", level=2)
        leave_demands = demands["leaves"]

        items = []
        if leave_demands.get("sick_children_days"):
            items.append(
                f"{leave_demands['sick_children_days']} jours pour enfants malades"
            )
        if leave_demands.get("caregiver_days"):
            items.append(f"{leave_demands['caregiver_days']} jours aidant")

        add_bullet_list(doc, items)
        add_styled_paragraph(
            doc,
            f"Base légale : {format_legal_reference('L1225-65')}"
        )
        doc.add_paragraph()

    # 3. Protection sociale
    if demands.get("health_insurance", {}).get("enabled"):
        add_heading(doc, "3. Protection Sociale", level=2)
        health_demands = demands["health_insurance"]

        add_styled_paragraph(
            doc,
            f"Mutuelle d'entreprise : prise en charge employeur à {health_demands.get('employer_contribution', 0)}%"
        )
        doc.add_paragraph()

    # 4. Conditions de travail
    if demands.get("working_conditions", {}).get("enabled"):
        add_heading(doc, "4. Conditions de Travail", level=2)
        wc_demands = demands["working_conditions"]

        items = []
        if wc_demands.get("prevention_plan"):
            items.append("Plan de prévention TMS et RPS")
        if wc_demands.get("equipment_renewal"):
            items.append("Renouvellement des équipements de travail")
        if wc_demands.get("shift_split_removal"):
            items.append("Suppression des coupures")

        add_bullet_list(doc, items)
        add_styled_paragraph(
            doc,
            f"Base légale : {format_legal_reference('L4121-1')}"
        )
        doc.add_paragraph()

    # Conclusion
    add_heading(doc, "Conclusion", level=2)
    add_styled_paragraph(
        doc,
        "Ces revendications visent à améliorer les conditions de travail et de vie "
        "des salarié·es tout en respectant le cadre légal des NAO. Nous restons ouverts "
        "au dialogue et à la négociation."
    )

    save_document(doc, output_path)
    return output_path


def generate_tract(campaign: Dict, output_path: str) -> str:
    """
    Generate tract for employees.

    Args:
        campaign: Campaign data dictionary
        output_path: Output file path

    Returns:
        Path to generated document
    """
    doc = create_document()

    # Title
    add_heading(
        doc,
        f"NAO {datetime.now().year} : Mobilisons-nous pour nos revendications !",
        level=1,
    )
    doc.add_paragraph()

    # Introduction
    add_styled_paragraph(
        doc,
        "Chers·ères collègues,",
        bold=True,
    )
    doc.add_paragraph()

    add_styled_paragraph(
        doc,
        f"Les Négociations Annuelles Obligatoires (NAO) {datetime.now().year} ont débuté. "
        f"Votre syndicat porte vos revendications auprès de la Direction."
    )
    doc.add_paragraph()

    # Main demands summary
    add_heading(doc, "Nos revendications principales :", level=2)

    demands = campaign.get("demands", {})
    demand_items = []

    if demands.get("salaries", {}).get("enabled"):
        increase = demands["salaries"].get("general_increase", "X")
        demand_items.append(f"✊ Augmentation des salaires de {increase}%")

    if demands.get("leaves", {}).get("enabled"):
        demand_items.append("✊ Plus de congés (enfants malades, aidant)")

    if demands.get("health_insurance", {}).get("enabled"):
        demand_items.append("✊ Meilleure mutuelle d'entreprise")

    if demands.get("working_conditions", {}).get("enabled"):
        demand_items.append("✊ Amélioration des conditions de travail")

    for item in demand_items:
        add_styled_paragraph(doc, item, bold=True)

    doc.add_paragraph()

    # Call to action
    add_heading(doc, "Ensemble, nous sommes plus forts !", level=2)
    add_styled_paragraph(
        doc,
        "Pour que ces revendications aboutissent, nous avons besoin de votre soutien :"
    )

    actions = [
        "Signez la pétition de soutien",
        "Parlez-en autour de vous",
        "Participez aux Assemblées Générales",
        "Contactez vos délégué·es syndicaux",
    ]
    add_bullet_list(doc, actions)
    doc.add_paragraph()

    # Contact
    add_styled_paragraph(
        doc,
        f"Contact : {campaign.get('organization_name', 'CGT')}",
        bold=True,
    )

    save_document(doc, output_path)
    return output_path


def generate_all_documents(campaign: Dict, output_dir: str) -> List[str]:
    """
    Generate all NAO documents for a campaign.

    Args:
        campaign: Campaign data dictionary
        output_dir: Output directory path

    Returns:
        List of generated document paths
    """
    os.makedirs(output_dir, exist_ok=True)

    generated_files = []

    # Generate letter
    letter_path = os.path.join(output_dir, "lettre_ouverture.docx")
    generate_letter_ouverture(campaign, letter_path)
    generated_files.append(letter_path)

    # Generate demands
    demands_path = os.path.join(output_dir, "revendications.docx")
    generate_revendications(campaign, demands_path)
    generated_files.append(demands_path)

    # Generate tract
    tract_path = os.path.join(output_dir, "tract.docx")
    generate_tract(campaign, tract_path)
    generated_files.append(tract_path)

    return generated_files
