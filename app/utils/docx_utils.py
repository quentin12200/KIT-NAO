"""Utilities for python-docx document generation."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import List, Tuple


def create_document() -> Document:
    """Create a new Word document with default settings."""
    doc = Document()
    return doc


def set_document_margins(
    doc: Document,
    top: float = 1.0,
    bottom: float = 1.0,
    left: float = 1.0,
    right: float = 1.0,
) -> None:
    """Set document margins (in inches)."""
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(top)
        section.bottom_margin = Inches(bottom)
        section.left_margin = Inches(left)
        section.right_margin = Inches(right)


def add_styled_paragraph(
    doc: Document, text: str, style: str = "Normal", bold: bool = False
) -> None:
    """Add a paragraph with specified style."""
    paragraph = doc.add_paragraph(text, style=style)
    if bold:
        for run in paragraph.runs:
            run.bold = True


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    """Add a heading to the document."""
    doc.add_heading(text, level=level)


def add_table_with_data(
    doc: Document, headers: List[str], rows: List[List[str]]
) -> None:
    """Add a table with headers and data."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"

    # Add headers
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        # Make header bold
        for paragraph in header_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True

    # Add data rows
    for row_idx, row_data in enumerate(rows, start=1):
        row_cells = table.rows[row_idx].cells
        for col_idx, cell_data in enumerate(row_data):
            row_cells[col_idx].text = str(cell_data)


def add_bullet_list(doc: Document, items: List[str]) -> None:
    """Add a bulleted list to the document."""
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered_list(doc: Document, items: List[str]) -> None:
    """Add a numbered list to the document."""
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_signature_block(
    doc: Document, name: str, title: str, organization: str
) -> None:
    """Add a signature block to the document."""
    doc.add_paragraph()  # Empty line
    doc.add_paragraph(f"Fait à __________, le __________")
    doc.add_paragraph()  # Empty line
    doc.add_paragraph(name)
    doc.add_paragraph(title)
    doc.add_paragraph(organization)


def save_document(doc: Document, filepath: str) -> None:
    """Save the document to a file."""
    doc.save(filepath)
